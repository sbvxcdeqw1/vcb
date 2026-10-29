"""Сверка с образцами заказчика: Кировоградская 180, кв. 24, расчёт на 13.04.2026."""
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

from vcb.calc.ledger import build_charge_lines, select_lines, total_debt
from vcb.calc.penalty import PenaltyRules, RateTable, calc_penalty, total_penalty
from vcb.models import SERVICES
from vcb.parsers.onec import parse_card

SAMPLE = Path(__file__).parent.parent / "examples" / "karta_kv24.xls"


@pytest.mark.skipif(not SAMPLE.exists(), reason="нет примера выгрузки")
def test_kv24_matches_sample():
    acc = parse_card(SAMPLE)
    assert acc.number == "00-0000000024"
    entries = [e for e in acc.entries if e.category == SERVICES]
    as_of = date(2026, 4, 13)
    lines = select_lines(build_charge_lines(entries, as_of=as_of), period_from=date(2024, 6, 1))
    assert total_debt(lines) == Decimal("24805.47")
    items = calc_penalty(lines, as_of, PenaltyRules(), RateTable(Decimal("9.5")))
    assert total_penalty(items) == Decimal("4351.96")


def test_statutory_mode_and_payments():
    """Оплата гасит самое старое начисление; пени по погашенной части идут до дня оплаты."""
    from vcb.models import CHARGE, PAYMENT, Entry

    E = lambda d, k, a: Entry(d, k, SERVICES, "x", Decimal(a))
    entries = [E(date(2025, 1, 31), CHARGE, "1000"), E(date(2025, 2, 28), CHARGE, "1000"),
               E(date(2025, 4, 1), PAYMENT, "1500")]
    as_of = date(2025, 6, 30)
    lines = select_lines(build_charge_lines(entries, as_of=as_of), period_from=date(2025, 1, 1))
    assert lines[0].debt == 0 and lines[1].debt == Decimal("500")
    assert total_debt(lines) == Decimal("500.00")
    items = calc_penalty(lines, as_of, PenaltyRules(mode="statutory"), RateTable(Decimal("10")))
    # январь: срок 10.02, пени с 13.03 (31-й день) по 31.03 включительно = 19 дней по 1/300
    jan = [s for i in items if i.line is lines[0] for s in i.segments]
    assert len(jan) == 1 and jan[0].days == 19 and jan[0].divisor == 300


def test_court_lookup():
    from vcb.court.registry import Court, CourtRegistry, house_matches
    from vcb.models import Debtor

    assert house_matches("2-40 чёт", "12") and not house_matches("2-40 чёт", "13")
    reg = CourtRegistry([Court("a", "magistrate", "Участок 1", territory=[
        {"city": "Пермь", "streets": [{"street": "ул. Кировоградская", "houses": "100-200"}]}])])
    assert reg.find(Debtor(city="Пермь", street="Кировоградская", house="180"), "magistrate").name == "Участок 1"
    assert reg.find(Debtor(city="Пермь", street="Кировоградская", house="20"), "magistrate") is None


CHESS = Path(__file__).parent.parent / "examples" / "shakhmatka_kv98.xlsx"


@pytest.mark.skipif(not CHESS.exists(), reason="нет примера шахматки")
def test_chess_parse_and_full_debt_matches_closing_balance():
    from vcb.models import OPENING
    from vcb.parsers.onec import parse_any

    acc = parse_any(CHESS)
    assert acc.number == "551010009822" and not acc.warnings
    lines = build_charge_lines(acc.entries, as_of=date(2026, 7, 31), with_opening=True)  # с начальным остатком
    assert total_debt(lines) == Decimal("79793.41")  # «Конечный остаток» образца
    assert sum(e.amount for e in acc.entries if e.kind == OPENING) == Decimal("5859.31")


CARD98 = Path(__file__).parent.parent / "examples" / "karta_kv98.xls"


@pytest.mark.skipif(not (CHESS.exists() and CARD98.exists()), reason="нет примеров кв. 98")
def test_card_vs_chess_reconciliation():
    from vcb.parsers.onec import parse_any
    from vcb.reconcile import compare, reconciled

    card, chess = parse_any(CARD98), parse_any(CHESS)
    assert card.saldo_start == Decimal("6060.19") and card.saldo_end == Decimal("162748.26")
    rep = compare(card, chess)
    # «Итог:» начального остатка в образце шахматки (6 935,03) - сальдо карточки с пенями на 31.05.2020
    assert rep.card_balance_start == Decimal("6935.03")
    assert rep.opening_diff == Decimal("1075.72")
    assert rep.principal_gap == Decimal("11269.54")  # в шахматке нет части услуг
    merged = reconciled(card, chess)
    by = lambda acc, k: sum(e.amount for e in acc.entries if e.kind == k and e.date >= date(2020, 6, 1))
    assert by(merged, "charge") == by(card, "charge") and by(merged, "payment") == by(card, "payment")


def test_payment_classification_and_overrides(tmp_path):
    from vcb.config import load_config
    from vcb.parsers.payments import ENFORCEMENT, RESIDENT, THIRD_PARTY, Payment, apply_overrides, classify

    mk = lambda i, d, a, t: Payment(f"П{i}", d, Decimal(a), t)
    ps = [mk(1, date(2019, 6, 24), "0.62", "Погашение долга Судебный приказ N2-1968/19 от 16.04.2019, 37522/19/59003-ИП"),
          mk(2, date(2020, 1, 20), "3000", "551010009822 12.2019-12.2019 к/плата=3000.00"),
          mk(3, date(2025, 6, 5), "978.63", "(991-0501-1540121500-244,л/с02991018855,225) ПД за 01.22")]
    classify(ps, load_config())
    assert [p.kind for p in ps] == [ENFORCEMENT, RESIDENT, THIRD_PARTY]
    assert [p.include for p in ps] == [False, True, True] and not ps[2].flags  # третье лицо зачитывается
    ov = tmp_path / "ov.csv"
    ov.write_text("дата;сумма;решение;комментарий\n05.06.2025;978.63;исключить;ИП\n", encoding="utf-8")
    apply_overrides(ps, ov)
    assert ps[2].include is False


def test_payments_go_to_oldest_charges_regardless_of_period_in_purpose():
    from vcb.models import CHARGE, PAYMENT, Entry

    E = lambda d, k, a, ref="": Entry(d, k, SERVICES, "x", Decimal(a), ref=ref)
    entries = [E(date(2025, 1, 31), CHARGE, "100"), E(date(2025, 2, 28), CHARGE, "100"),
               E(date(2025, 3, 31), CHARGE, "100"), E(date(2025, 4, 15), PAYMENT, "150", "П1")]
    lines = build_charge_lines(entries, as_of=date(2025, 5, 1))
    assert [ln.debt for ln in lines] == [Decimal(0), Decimal(50), Decimal(100)]
    assert {p.ref for ln in lines for p in ln.portions if p.closed} == {"П1"}


def test_payment_goes_to_period_from_purpose_and_unmatched_to_oldest():
    from vcb.models import CHARGE, PAYMENT, Entry

    E = lambda d, k, a, ref="", per=None: Entry(d, k, SERVICES, "x", Decimal(a), ref=ref, period=per)
    base = [E(date(2025, 1, 31), CHARGE, "100"), E(date(2025, 2, 28), CHARGE, "100"), E(date(2025, 3, 31), CHARGE, "100")]
    ignored = []
    pays = [E(date(2025, 4, 15), PAYMENT, "130", "П1", (2025, 3)),   # март 100 + остаток 30 на январь
            E(date(2025, 4, 16), PAYMENT, "40", "П2", (2024, 12))]   # период до начала расчёта - не учитывается
    lines = build_charge_lines(base + pays, as_of=date(2025, 5, 1), ignored=ignored)
    assert [ln.debt for ln in lines] == [Decimal(70), Decimal(100), Decimal(0)]
    assert [i[0] for i in ignored] == ["П2"]
    # режим «на самые ранние»: период в назначении игнорируется
    lines = build_charge_lines(base + pays[:1], as_of=date(2025, 5, 1), by_purpose=False)
    assert [ln.debt for ln in lines] == [Decimal(0), Decimal(70), Decimal(100)]


def test_payment_for_not_yet_accrued_month_waits_for_accrual():
    from vcb.models import CHARGE, PAYMENT, Entry

    E = lambda d, k, a, ref="", per=None: Entry(d, k, SERVICES, "x", Decimal(a), ref=ref, period=per)
    entries = [E(date(2025, 1, 31), CHARGE, "100"), E(date(2025, 2, 28), CHARGE, "100"),
               E(date(2025, 2, 10), PAYMENT, "100", "П1", (2025, 2))]  # оплата за февраль до его начисления
    lines = build_charge_lines(entries, as_of=date(2025, 3, 1))
    assert [ln.debt for ln in lines] == [Decimal(100), Decimal(0)]


def test_period_is_strictly_operator_defined():
    from vcb.models import CHARGE, PAYMENT, Entry
    from vcb.pipeline import _in_period

    E = lambda d, k, per=None: Entry(d, k, SERVICES, "x", Decimal(1), period=per)
    a, b = date(2024, 6, 1), date(2025, 5, 31)
    assert _in_period(E(date(2024, 6, 30), CHARGE), a, b) and not _in_period(E(date(2024, 5, 31), CHARGE), a, b)
    assert not _in_period(E(date(2025, 6, 30), CHARGE), a, b)
    assert _in_period(E(date(2025, 8, 1), PAYMENT, (2025, 5)), a, b)  # за май 2025, оплачено позже
    assert not _in_period(E(date(2025, 1, 1), PAYMENT, (2024, 5)), a, b)  # за период до начала
    assert not _in_period(E(date(2024, 5, 1), PAYMENT), a, b) and _in_period(E(date(2024, 7, 1), PAYMENT), a, b)


def test_real_perm_register_matches_sample_court():
    from vcb.court.registry import CourtRegistry
    from vcb.models import Debtor

    reg = CourtRegistry.load("perm_krai")
    d = Debtor(city="Пермь", street="Кировоградская", house="180", flat="24")
    ms = reg.find(d, "magistrate")
    assert ms and "№ 6 Кировского района" in ms.name  # дом 180 входит в участок № 6 (ранее - № 2)
    dc = reg.find(d, "district")
    assert dc.name == "Кировский районный суд г. Перми" and dc.address == "614113, г. Пермь, ул. Кировоградская, д. 37"
    # населённый пункт без разобранных улиц - по тексту описания границ
    assert reg.find(Debtor(city="Александровск", street="Ленина", house="5"), "district").name.startswith("Александровский")


def test_rate_cap_and_history():
    from vcb.config import load_config, rules_from

    rules, rates = rules_from(load_config())
    assert rates.at(date(2026, 4, 13)) == Decimal("9.5")  # ключевая выше, но действует потолок 9,5%
    assert rates.at(date(2021, 6, 1)) == Decimal("5.00")  # до потолка - фактическая ставка ЦБ
    assert len(rules.moratoria) == 2


@pytest.mark.skipif(not SAMPLE.exists(), reason="нет примера выгрузки")
def test_full_pipeline_reproduces_sample_with_default_config():
    from vcb.config import load_config
    from vcb.models import Party
    from vcb.pipeline import build_case

    prof = {"claimant": Party(), "representative": Party(), "legal": {}}
    c = build_case(SAMPLE, None, prof, load_config(), date(2026, 4, 13), None)
    res = c.results[SERVICES]
    assert (res.debt, res.penalty) == (Decimal("24805.47"), Decimal("4351.96"))


def test_gui_builds_cli_arguments():
    pytest.importorskip("tkinter")
    import vcb_gui

    argv = vcb_gui.build_argv({"source": "a.xls", "out": "o", "region": "perm_krai", "payments": "p.xlsx",
                               "as_of": "13.04.2026", "period_from": "01.10.2023", "doc": "claim"})
    assert argv[:2] == ["a.xls", "-o"] and "--no-interactive" in argv
    assert argv[argv.index("--as-of") + 1] == "2026-04-13"
    assert argv[argv.index("--period-from") + 1] == "2023-10-01" and "--period-to" not in argv
    assert argv[argv.index("--payments") + 1] == "p.xlsx" and argv[argv.index("--doc") + 1] == "claim"
    with pytest.raises(ValueError):
        vcb_gui.build_argv({"source": ""})
    with pytest.raises(ValueError):
        vcb_gui.build_argv({"source": "a", "as_of": "31/12"})


def test_version_is_consistent():
    import re

    import vcb

    pyproject = (Path(__file__).parent.parent / "pyproject.toml").read_text(encoding="utf-8")
    assert re.search(rf'^version = "{re.escape(vcb.__version__)}"', pyproject, flags=re.M)
    changelog = (Path(__file__).parent.parent / "CHANGELOG.md").read_text(encoding="utf-8")
    assert f"## {vcb.__version__} " in changelog


def test_typography_and_print_setup(tmp_path):
    from openpyxl import load_workbook

    from vcb.config import load_config
    from vcb.docs.claim import NBSP, typo
    from vcb.models import Party
    from vcb.pipeline import build_case, write_case

    assert typo("24 805,47 руб. по ст. 155 № 5, д. 180") == f"24{NBSP}805,47{NBSP}руб. по ст.{NBSP}155 №{NBSP}5, д.{NBSP}180"
    if not SAMPLE.exists():
        return
    prof = {"claimant": Party(), "representative": Party(), "legal": {}}
    case = build_case(SAMPLE, None, prof, load_config(), date(2026, 4, 13), None)
    write_case(case, tmp_path)
    for name, landscape in (("Карточка_расчета_ЖКУ.xlsx", False), ("Расчет_пени_ЖКУ.xlsx", True), ("Шахматка.xlsx", True)):
        ws = load_workbook(tmp_path / name).worksheets[0]
        assert ws.page_setup.paperSize == 9 and (ws.page_setup.orientation == "landscape") == landscape
        assert ws.print_title_rows
    import docx

    doc = docx.Document(next(tmp_path.glob("*.docx")))
    sec = doc.sections[0]
    assert round(sec.left_margin.cm, 1) == 3.0 and round(sec.right_margin.cm, 1) == 1.0
    assert doc.styles["Normal"].font.size.pt == 14 and doc.styles["Normal"].font.name == "Times New Roman"


def test_every_xlsx_has_signature_block_with_representative(tmp_path):
    from openpyxl import load_workbook

    from vcb.config import load_config
    from vcb.models import Party
    from vcb.pipeline import build_case, write_case

    if not SAMPLE.exists():
        pytest.skip("нет примера выгрузки")
    prof = {"claimant": Party(short_name="ООО «Тест»"), "representative": Party(name="Петров Пётр Петрович",
                                                                            sign_name="П.П.Петров"), "legal": {}}
    case = build_case(SAMPLE, None, prof, load_config(), date(2026, 4, 13), None)
    write_case(case, tmp_path)
    files = sorted(tmp_path.glob("*.xlsx"))
    assert len(files) >= 3
    for f in files:
        for ws in load_workbook(f).worksheets:
            text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
            assert "Представитель ООО «Тест»" in text and "/ П.П.Петров /" in text and "(Петров Пётр Петрович)" in text, \
                f"{f.name}/{ws.title}"
