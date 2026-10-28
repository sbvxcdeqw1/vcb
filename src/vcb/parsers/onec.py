"""Разбор выгрузки 1С «Карточка расчетов» (xls/xlsx)."""
from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from pathlib import Path

from ..models import CAPREPAIR, CHARGE, OPENING, PAYMENT, PENALTY, SERVICES, Account, Entry
from ..money import D

CAPREPAIR_RE = re.compile(r"капитальн|капремонт|кап\.\s*ремонт", re.I)
LS_RE = re.compile(r"по лицевому счету\s*([\w\-]+)\s*,?\s*(.*)", re.I)
DATE_RE = re.compile(r"^\d{2}\.\d{2}\.\d{4}$")


class ParseError(ValueError):
    pass


def _rows(path: Path) -> list[list]:
    if path.suffix.lower() == ".xls":
        import xlrd

        wb = xlrd.open_workbook(str(path))
        sh = wb.sheet_by_index(0)
        out = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, wb.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                else:
                    row.append(cell.value)
            out.append(row)
        return out
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and DATE_RE.match(v.strip()):
        return datetime.strptime(v.strip(), "%d.%m.%Y").date()
    return None


def _num(v):
    if v in (None, ""):
        return None
    return D(v)


def classify_category(*texts: str) -> str:
    return CAPREPAIR if CAPREPAIR_RE.search(" ".join(texts)) else SERVICES


def parse_any(path: str | Path) -> Account:
    """«Шахматка» (по услугам) или «Карточка расчетов» - определяем по содержимому."""
    from .chess import is_chess, parse_chess

    path = Path(path)
    rows = _rows(path)
    return parse_chess(path, rows) if is_chess(rows) else parse_card(path, rows)


def parse_card(path: str | Path, rows: list[list] | None = None) -> Account:
    path = Path(path)
    rows = rows if rows is not None else _rows(path)
    acc = Account(source=path.name)
    head = None
    prev_running = None
    for i, row in enumerate(rows):
        cells = [str(c).strip() for c in row if c not in (None, "")]
        if not cells:
            continue
        if head is None:
            text = " ".join(cells)
            m = LS_RE.search(text)
            if m:
                acc.number, acc.owner = m.group(1), m.group(2).strip()
                # адрес - следующая непустая строка
                for nxt in rows[i + 1:i + 3]:
                    t = " ".join(str(c).strip() for c in nxt if c not in (None, ""))
                    if t and not t.lower().startswith("дата"):
                        acc.address = t
                        break
            if cells[0].lower() == "дата" and any("регистратор" in c.lower() for c in cells):
                head = {c.strip().lower(): j for j, c in enumerate(row) if isinstance(c, str)}
            continue
        d = _date(row[head["дата"]])
        if d is None:
            label = str(row[0] or "")
            m = re.match(r"Сальдо на (\d{2}\.\d{2}\.\d{4})", label)
            bal = _num(row[head["задолженность"]]) if "задолженность" in head else None
            if m and bal is not None:
                if acc.saldo_start is None:
                    acc.saldo_start, acc.saldo_start_date = bal, datetime.strptime(m.group(1), "%d.%m.%Y").date()
                else:
                    acc.saldo_end = bal
            continue  # «Обороты» и пр.
        running = _num(row[head["задолженность"]]) if "задолженность" in head else None
        if running is not None and acc.saldo_start is not None:
            prev = prev_running if prev_running is not None else acc.saldo_start
            delta = (_num(row[head["начислено"]]) or 0) - (_num(row[head["оплата"]]) or 0)
            if abs(prev + delta - running) > D("0.005"):
                acc.warnings.append(f"{d:%d.%m.%Y}: бегущее сальдо {running} не сходится с {prev + delta}")
            prev_running = running
        registrar = str(row[head["регистратор"]] or "")
        oper = str(row[head.get("операция", head["регистратор"])] or "")
        charged = _num(row[head["начислено"]]) if "начислено" in head else None
        paid = _num(row[head["оплата"]]) if "оплата" in head else None
        cat = classify_category(oper, registrar)
        if "пени" in oper.lower() or "пени" in registrar.lower():
            kind, amount = PENALTY, charged or paid
        elif paid:
            kind, amount = PAYMENT, paid
        elif charged is not None:
            kind, amount = CHARGE, charged
        else:
            continue
        if amount is None:
            continue
        acc.entries.append(Entry(d, kind, cat, oper.strip(), amount, registrar.strip()))
    if head is None:
        raise ParseError(f"{path.name}: не найдена таблица с колонками Дата/Регистратор/Операция")
    if acc.saldo_start and acc.saldo_start > 0:
        # начальное сальдо карточки включает и пени - разбить его по карточке нельзя
        acc.entries.append(Entry(acc.saldo_start_date - timedelta(days=1), OPENING, SERVICES,
                                 f"Сальдо на {acc.saldo_start_date:%d.%m.%Y}", acc.saldo_start))
    return acc
