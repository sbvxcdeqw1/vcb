"""Карточка расчёта задолженности и расчёт пени (xlsx) - по образцам заказчика, оформление для печати на A4."""
from __future__ import annotations

from decimal import Decimal

from openpyxl import Workbook

from .. import __version__
from ..models import CAPREPAIR, CATEGORY_TITLES, CHARGE, OPENING, PAYMENT
from ..money import fmt as money_fmt
from .context import Case, CategoryResult
from .xlsx_util import (LEFT, MONEY, TOTAL_FILL, header_row, put, setup_print, signature_block, title_block,
                        widths)


def _lines(case: Case, title: str, category: str) -> list[tuple[str, str]]:
    org = case.claimant.short_name or case.claimant.name or "Взыскатель: ____________________"
    what = (f"по услугам: {CATEGORY_TITLES[category]}" if category != CAPREPAIR
            else "по взносам на капитальный ремонт")
    lines = [(org, "org"), (title.upper(), "title"), (what, "center"),
             (f"за период с {case.period_from:%d.%m.%Y} по {case.period_to:%d.%m.%Y}, "
              f"по состоянию на {case.as_of:%d.%m.%Y}", "center"), ("", "gap"),
             (f"Адрес: {case.debtor.address or case.account.address}", "left")]
    if case.debtor.fio or case.account.owner:
        lines.append((f"Собственник: {case.debtor.fio or case.account.owner}", "left"))
    if case.debtor.area:
        lines.append((f"Общая площадь квартиры: {case.debtor.area} кв. м", "left"))
    lines.append(("", "gap"))
    return lines


def write_calc_card(path, case: Case, res: CategoryResult):
    """Карточка расчёта задолженности (по образцам «Расчет задолженности собственника жилья» и «Карточка расчетов»):
    только тело долга и оплаты, без пеней. Начисления - из шахматки, оплаты - из реестра, нарастающий итог по строкам."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Карточка расчёта"
    lines = _lines(case, "Расчёт задолженности собственника жилья", res.category)
    hdr = title_block(ws, 6, lines)
    header_row(ws, hdr, ["Дата", "Регистратор", "Операция", "Начислено, руб.", "Оплата, руб.", "Задолженность, руб."])
    r = hdr + 1
    start = r
    for when, registrar, purpose, charged, paid in res.card_rows:
        put(ws, f"A{r}", when, border=True, fmt="DD.MM.YYYY", align=LEFT)
        put(ws, f"B{r}", registrar, border=True, align=LEFT)
        put(ws, f"C{r}", purpose, border=True, align=LEFT)
        put(ws, f"D{r}", charged or None, border=True, fmt=MONEY)
        put(ws, f"E{r}", paid or None, border=True, fmt=MONEY)
        put(ws, f"F{r}", f"=F{r - 1}+D{r}-E{r}" if r > start else f"=D{r}-E{r}", border=True, fmt=MONEY)
        r += 1
    put(ws, f"A{r}", "Обороты за период", bold=True, border=True, fill=TOTAL_FILL)
    for col in "BC":
        put(ws, f"{col}{r}", None, border=True, fill=TOTAL_FILL)
    ws.merge_cells(f"A{r}:C{r}")
    put(ws, f"D{r}", f"=SUM(D{start}:D{r - 1})", bold=True, border=True, fmt=MONEY, fill=TOTAL_FILL)
    put(ws, f"E{r}", f"=SUM(E{start}:E{r - 1})", bold=True, border=True, fmt=MONEY, fill=TOTAL_FILL)
    put(ws, f"F{r}", None, border=True, fill=TOTAL_FILL)
    put(ws, f"A{r + 1}", f"Задолженность на {case.as_of:%d.%m.%Y}", bold=True, border=True, fill=TOTAL_FILL)
    for col in "BCDE":
        put(ws, f"{col}{r + 1}", None, border=True, fill=TOTAL_FILL)
    ws.merge_cells(f"A{r + 1}:E{r + 1}")
    put(ws, f"F{r + 1}", f"=F{r - 1}" if r > start else 0, bold=True, border=True, fmt=MONEY, fill=TOTAL_FILL)
    nxt = r + 3
    signature_block(ws, nxt + 1, case, ncols=3)
    widths(ws, {"A": 12, "B": 36, "C": 24, "D": 15, "E": 15, "F": 17})
    setup_print(ws, landscape=False, title_rows=f"{hdr}:{hdr}", footer_left=f"vcb {__version__}")
    wb.save(path)


write_debt_card = write_calc_card  # прежнее имя


def _merge_items(res: CategoryResult) -> dict:
    """Части начисления, погашенные в один день (оплаты по разным услугам), показываем одной строкой расчёта."""
    from ..calc.penalty import PenaltyItem, Segment

    out: dict[int, list] = {}
    for it in res.penalty_items:
        out.setdefault(id(it.line), []).append(it)
    for key, items in out.items():
        by_date: dict = {}
        for it in items:
            by_date.setdefault(it.pay_date, []).append(it)
        merged = []
        for pay_date, group in by_date.items():
            if len(group) == 1:
                merged.append(group[0])
                continue
            segs: dict[tuple, Segment] = {}
            for it in group:
                for sg in it.segments:
                    k = (sg.start, sg.end, sg.days, sg.divisor, sg.rate)
                    if k in segs:
                        segs[k].base += sg.base
                        segs[k].penalty += sg.penalty
                    else:
                        segs[k] = Segment(sg.start, sg.end, sg.days, sg.divisor, sg.rate, sg.base, sg.penalty)
            merged.append(PenaltyItem(group[0].line, sum((i.base for i in group), Decimal(0)), pay_date,
                                      list(segs.values())))
        out[key] = merged
    return out


def write_penalty_card(path, case: Case, res: CategoryResult):
    """Образец: «Расчет пени по задолженности собственника жилья»."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Расчет пени"
    hdr = title_block(ws, 12, _lines(case, "Расчёт пени по задолженности собственника жилья", res.category))
    header_row(ws, hdr, ["Месяц", "Начислено за месяц, руб.", "Подлежит оплате до", "Дата оплаты", "Сумма оплаты, руб.",
                         "Ставка ЦБ, %", "Делитель", "Доля ставки", "Расчётная сумма, руб.", "Период просрочки",
                         "Количество дней", "Сумма пени, руб."], height=44)
    r = hdr + 1
    first_row = r
    items_by_line = _merge_items(res)
    for ln in res.lines:
        first = True
        for it in items_by_line.get(id(ln), []):
            for k, s in enumerate(it.segments or [None]):
                put(ws, f"A{r}", ln.label if first else None, border=True, align=LEFT)
                put(ws, f"B{r}", ln.amount if first else None, border=True, fmt=MONEY)
                put(ws, f"C{r}", f"{ln.due:%d.%m.%Y}" if first else None, border=True, align=LEFT)
                put(ws, f"D{r}", f"{it.pay_date:%d.%m.%Y}" if it.pay_date and k == 0 else None, border=True,
                    align=LEFT)
                put(ws, f"E{r}", it.base if it.pay_date and k == 0 else None, border=True, fmt=MONEY)
                put(ws, f"I{r}", it.base, border=True, fmt=MONEY)
                if s:
                    put(ws, f"F{r}", s.rate, border=True, fmt="0.00")
                    put(ws, f"G{r}", s.divisor, border=True, fmt="0")
                    put(ws, f"H{r}", f"=F{r}/G{r}", border=True, fmt="0.0000")
                    put(ws, f"J{r}", s.period, border=True, align=LEFT)
                    put(ws, f"K{r}", s.days, border=True, fmt="0")
                    put(ws, f"L{r}", f"=I{r}*H{r}*K{r}/100", border=True, fmt=MONEY)
                else:
                    for col in "FGHJKL":
                        put(ws, f"{col}{r}", None, border=True)
                first = False
                r += 1
        if first:  # начисление без строк расчёта (пени ещё не начисляются)
            put(ws, f"A{r}", ln.label, border=True, align=LEFT)
            put(ws, f"B{r}", ln.amount, border=True, fmt=MONEY)
            put(ws, f"C{r}", f"{ln.due:%d.%m.%Y}", border=True, align=LEFT)
            for col in "DEFGHIJKL":
                put(ws, f"{col}{r}", None, border=True)
            r += 1
    put(ws, f"A{r}", "Итого", bold=True, border=True, fill=TOTAL_FILL)
    for col in "CDFGHIJK":
        put(ws, f"{col}{r}", None, border=True, fill=TOTAL_FILL)
    # начисления учитываем один раз - по строкам, где заполнена колонка C (первая строка месяца)
    put(ws, f"B{r}", f'=SUMIF(C{first_row}:C{r - 1},"<>",B{first_row}:B{r - 1})', bold=True, border=True, fmt=MONEY,
        fill=TOTAL_FILL)
    put(ws, f"E{r}", f"=SUM(E{first_row}:E{r - 1})", bold=True, border=True, fmt=MONEY, fill=TOTAL_FILL)
    put(ws, f"L{r}", f"=SUM(L{first_row}:L{r - 1})", bold=True, border=True, fmt=MONEY, fill=TOTAL_FILL)
    put(ws, f"A{r + 2}", f"Всего: {money_fmt(res.debt)} руб. (задолженность)", bold=True)
    put(ws, f"A{r + 3}", f"            {money_fmt(res.penalty)} руб. (пени)", bold=True)
    signature_block(ws, r + 6, case, ncols=3)
    widths(ws, {"A": 17, "B": 14, "C": 13, "D": 12, "E": 14, "F": 8, "G": 10, "H": 9, "I": 14, "J": 19, "K": 12,
                "L": 14})
    setup_print(ws, landscape=True, title_rows=f"{hdr}:{hdr}", footer_left=f"vcb {__version__}")
    wb.save(path)
