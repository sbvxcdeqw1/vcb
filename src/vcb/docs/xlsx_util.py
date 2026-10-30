"""Оформление xlsx: единый шрифт, шапки, печать на A4 (поля, повтор заголовков, нумерация страниц)."""
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break

FONT_NAME = "Arial"
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="E3E8EF")
TOTAL_FILL = PatternFill("solid", fgColor="F2F4F7")
NORMAL = Font(name=FONT_NAME, size=10)
BOLD = Font(name=FONT_NAME, size=10, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
MONEY = "#,##0.00"


def cap(text):
    """Название с заглавной буквы: «содержание жилья» -> «Содержание жилья»."""
    text = str(text or "").strip()
    return text[:1].upper() + text[1:]


def put(ws, ref, value, *, bold=False, border=False, align=None, fmt=None, size=10, fill=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT_NAME, size=size, bold=bold)
    if border:
        c.border = BORDER
    c.alignment = align or (RIGHT if fmt == MONEY else Alignment(vertical="center"))
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def widths(ws, spec: dict):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def title_block(ws, ncols: int, lines: list[tuple[str, str]], width_chars: float | None = None) -> int:
    """Шапка документа над таблицей. Стили строк: org - организация (жирный, слева), title - заголовок (по центру,
    крупно), center - по центру, left - слева, gap - пустая строка. Длинные строки переносятся, высота подбирается
    по ширине блока (`width_chars`, по умолчанию 14 знаков на столбец). Возвращает первую свободную строку."""
    import math

    last = ws.cell(1, ncols).column_letter
    width = width_chars or 14 * ncols
    r = 1
    for text, style in lines:
        if style == "gap":
            ws.row_dimensions[r].height = 8
        else:
            ws.merge_cells(f"A{r}:{last}{r}")
            big = style == "title"
            put(ws, f"A{r}", text, bold=style in ("org", "title"), size=12 if big else 10,
                align=CENTER if style in ("title", "center") else LEFT)
            n = max(1, math.ceil(len(text) * (1.3 if big else 1.0) / width))
            ws.row_dimensions[r].height = (18 if big else 13.5) * n + 3
        r += 1
    return r


def header_row(ws, row: int, labels: list[str], height: float = 32):
    for i, text in enumerate(labels, 1):
        put(ws, ws.cell(row, i).coordinate, text, bold=True, border=True, align=CENTER, fill=HEADER_FILL)
    ws.row_dimensions[row].height = height


def setup_print(ws, *, landscape: bool, title_rows: str | None = None, title_cols: str | None = None,
                fit_width: bool = True, scale: int | None = None, col_breaks: tuple = (), footer_left: str = ""):
    """A4, поля 1,5/1/1,5/1,8 см, повтор шапки таблицы на каждой странице, номера страниц «Страница N из M»."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    if fit_width:
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0  # по высоте - сколько потребуется страниц
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    elif scale:
        ws.page_setup.scale = scale
    ws.page_margins = PageMargins(left=0.59, right=0.39, top=0.59, bottom=0.71, header=0.3, footer=0.3)
    ws.print_options.horizontalCentered = True
    if title_rows:
        ws.print_title_rows = title_rows
    if title_cols:
        ws.print_title_cols = title_cols
    for col in col_breaks:
        ws.col_breaks.append(Break(id=col))
    ws.oddFooter.left.text = footer_left
    ws.oddFooter.left.size = 8
    ws.oddFooter.center.text = "Страница &P из &N"
    ws.oddFooter.center.size = 8
    ws.oddFooter.right.text = "Напечатано &D"
    ws.oddFooter.right.size = 8


def printable(ws, landscape=True):
    """Совместимость: простая печать на A4 по ширине."""
    setup_print(ws, landscape=landscape)


def signature_block(ws, row: int, case, ncols: int = 3) -> int:
    """Поле подписи под таблицей: должность и организация, линия для подписи с расшифровкой (как задана в окне
    программы), дата. Блок занимает `ncols` левых столбцов. Возвращает следующую свободную строку."""
    org = case.claimant.short_name or case.claimant.name or "____________________"
    rep = case.representative
    who = "Представитель" if (rep.name or rep.sign_name) else "Взыскатель"
    sign = rep.sign_name or rep.name or case.claimant.sign_name or "____________"
    last = get_column_letter(ncols)
    rows = [(f"{who} {org}", Font(name=FONT_NAME, size=10, bold=True), 28),
            (f"_____________________ / {sign} /", NORMAL, 20),
            ("Дата: «____» ______________ 20____ г.", NORMAL, 18)]
    for i, (text, font, height) in enumerate(rows):
        r = row + i
        ws.merge_cells(f"A{r}:{last}{r}")
        c = ws[f"A{r}"]
        c.value, c.font = text, font
        c.alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True)
        ws.row_dimensions[r].height = height
    return row + len(rows)
