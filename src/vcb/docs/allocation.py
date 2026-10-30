"""Отчёт о разноске оплат: что зачтено, что исключено, на какие начисления перенесено."""
from __future__ import annotations

from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook

from ..money import fmt
from ..parsers.payments import KIND_TITLES
from .context import Case
from .. import __version__
from .xlsx_util import (BORDER, CENTER, HEADER_FILL, LEFT, MONEY, TOTAL_FILL, header_row, put, setup_print, signature_block,
                        title_block, widths)


def write_allocation(path, case: Case, res):
    skipped = {ref: why for ref, _, why in getattr(res, "ignored", [])}
    refs = defaultdict(list)  # ref -> [(начисление, сумма)]
    paid_by = defaultdict(list)  # id строки -> [(оплата, сумма)]
    for ln in res.all_lines:
        for p in ln.portions:
            if p.closed is not None and p.ref:
                refs[p.ref].append((ln.label, p.amount))
                paid_by[id(ln)].append((p.ref, p.closed, p.amount))
    wb = Workbook()
    ws = wb.active
    ws.title = "Оплаты"
    org = case.claimant.short_name or case.claimant.name
    hdr = title_block(ws, 10, ([(org, "org")] if org else []) + [
        ("РАЗНОСКА ОПЛАТ", "title"),
        (f"{case.debtor.fio or (res.account or case.account).owner}, лицевой счёт "
         f"{(res.account or case.account).number}", "center"),
        ("", "gap")], width_chars=150)
    header_row(ws, hdr, ["№", "Дата", "Сумма, руб.", "Тип платежа", "Зачтено", "Основание", "Период по назначению",
                         "Разнесено на", "Примечание", "Назначение платежа"], height=32)
    first_data = hdr + 1
    r = first_data
    for p in sorted(res.payments, key=lambda p: (p.date, p.ref)):
        parts = refs.get(p.ref, [])
        put(ws, f"A{r}", p.ref, border=True)
        put(ws, f"B{r}", p.date, border=True, fmt="DD.MM.YYYY")
        put(ws, f"C{r}", p.amount, border=True, fmt=MONEY)
        put(ws, f"D{r}", KIND_TITLES[p.kind], border=True, align=LEFT)
        put(ws, f"E{r}", "да" if p.include else "нет", bold=not p.include, border=True)
        put(ws, f"F{r}", p.reason, border=True, align=LEFT)
        put(ws, f"G{r}", f"{p.period_hint[1]:02d}.{p.period_hint[0]}" if p.period_hint else "", border=True)
        where = "; ".join(f"{lbl}: {fmt(a)}" for lbl, a in parts)
        if p.ref in skipped:
            where = f"не учтена: {skipped[p.ref]}"
        elif not where:
            where = "не учтена: вне расчётного периода"
        put(ws, f"H{r}", where if p.include else "—", border=True, align=LEFT)
        put(ws, f"I{r}", "; ".join(p.flags), border=True, align=LEFT)
        put(ws, f"J{r}", p.info[:160] + ("…" if len(p.info) > 160 else ""), border=True, align=LEFT)
        r += 1
    put(ws, f"A{r}", "Итого зачтено", bold=True, border=True)
    put(ws, f"C{r}", f'=SUMIF(E{first_data}:E{r - 1},"да",C{first_data}:C{r - 1})', bold=True, border=True, fmt=MONEY)
    put(ws, f"A{r + 1}", "Итого не зачтено", bold=True, border=True)
    put(ws, f"C{r + 1}", f'=SUMIF(E{first_data}:E{r - 1},"нет",C{first_data}:C{r - 1})', bold=True, border=True, fmt=MONEY)
    signature_block(ws, r + 4, case, ncols=5)
    widths(ws, {"A": 6, "B": 11, "C": 12, "D": 27, "E": 9, "F": 30, "G": 12, "H": 44, "I": 22, "J": 60})
    ws.freeze_panes = ws.cell(first_data, 1)
    setup_print(ws, landscape=True, title_rows=f"{hdr}:{hdr}", footer_left=f"vcb {__version__}")

    w2 = wb.create_sheet("Начисления")
    header_row(w2, 1, ["Начисление", "Срок оплаты", "Сумма, руб.", "Погашено оплатами", "Остаток долга, руб.",
                       "В заявлении"])
    included = {id(ln) for ln in res.lines}
    for r, ln in enumerate(res.all_lines, 2):
        put(w2, f"A{r}", ln.label, border=True)
        put(w2, f"B{r}", ln.due, border=True, fmt="DD.MM.YYYY")
        put(w2, f"C{r}", ln.amount, border=True, fmt=MONEY)
        put(w2, f"D{r}", "; ".join(f"{ref} {c:%d.%m.%Y} {fmt(a)}" for ref, c, a in paid_by.get(id(ln), [])),
            border=True, align=LEFT)
        put(w2, f"E{r}", ln.debt, border=True, fmt=MONEY)
        put(w2, f"F{r}", "да" if id(ln) in included else "нет (давность / вне периода / остаток)", border=True)
    signature_block(w2, len(res.all_lines) + 4, case, ncols=3)
    widths(w2, {"A": 20, "B": 13, "C": 12, "D": 70, "E": 13, "F": 34})
    w2.freeze_panes = "A2"
    setup_print(w2, landscape=True, title_rows="1:1", footer_left=f"vcb {__version__}")

    w3 = wb.create_sheet("Сценарии")
    header_row(w3, 1, ["Вариант учёта оплат", "Долг, руб.", "Пени, руб.", "Итого, руб."], height=24)
    for r, (name, debt, pen) in enumerate(res.scenarios, 2):
        put(w3, f"A{r}", name, border=True)
        put(w3, f"B{r}", debt, border=True, fmt=MONEY)
        put(w3, f"C{r}", pen, border=True, fmt=MONEY)
        put(w3, f"D{r}", f"=B{r}+C{r}", border=True, fmt=MONEY)
    n = len(res.scenarios) + 3
    put(w3, f"A{n}", "Принято в заявлении (с учётом настроек и ручных решений)", bold=True, border=True)
    put(w3, f"B{n}", res.debt, bold=True, border=True, fmt=MONEY)
    put(w3, f"C{n}", res.penalty, bold=True, border=True, fmt=MONEY)
    put(w3, f"D{n}", f"=B{n}+C{n}", bold=True, border=True, fmt=MONEY)
    signature_block(w3, n + 3, case, ncols=2)
    widths(w3, {"A": 62, "B": 15, "C": 15, "D": 15})
    setup_print(w3, landscape=False, footer_left=f"vcb {__version__}")
    wb.save(path)
